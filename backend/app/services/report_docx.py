"""Generate an editable Word document from the integrated teaching plan."""

from __future__ import annotations

from io import BytesIO


def render_report_docx(result: dict) -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as error:
        raise RuntimeError("Word 导出依赖未安装，请安装项目 requirements.txt 中的 python-docx") from error

    plan = result["integrated_plan"]
    basic = plan["basic_information"]
    diagnosis = plan["diagnosis_summary"]
    context = plan["teaching_context_and_conditions"]
    goals = plan["goal_design"]
    audit = result["audit_summary"]
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    def set_font(style, size: int, bold: bool = False):
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for part in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{part}"), "宋体" if part == "eastAsia" else "Arial")

    set_font(document.styles["Normal"], 10)
    set_font(document.styles["Title"], 18, True)
    set_font(document.styles["Heading 1"], 13, True)
    set_font(document.styles["Heading 2"], 11, True)
    title_border = OxmlElement("w:pBdr")
    bottom_border = OxmlElement("w:bottom")
    bottom_border.set(qn("w:val"), "nil")
    title_border.append(bottom_border)
    document.styles["Title"]._element.get_or_add_pPr().append(title_border)
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.35
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].paragraph_format.space_before = Pt(12)
        document.styles[name].paragraph_format.space_after = Pt(6)

    def value(item: object) -> str:
        if item is None or item == "":
            return "—"
        if isinstance(item, list):
            return "；".join(str(part) for part in item)
        return str(item)

    def field(label: str, item: object):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(value(item))

    def table(headers: list[str] | None, rows: list[list[object]], widths: list[float]):
        word_table = document.add_table(rows=1 if headers else 0, cols=len(widths))
        word_table.autofit = False
        for index, width in enumerate(widths):
            word_table.columns[index].width = Cm(width)
        if headers:
            for index, header in enumerate(headers):
                cell = word_table.rows[0].cells[index]
                cell.text = header
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EDF3F1")
                cell._tc.get_or_add_tcPr().append(shading)
            header_row = word_table.rows[0]._tr.get_or_add_trPr()
            header_row.append(OxmlElement("w:tblHeader"))
        for row in rows:
            cells = word_table.add_row().cells
            for index, item in enumerate(row):
                cells[index].text = value(item)
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for row in word_table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "D9D9D9")
            borders.append(border)
        word_table._tbl.tblPr.append(borders)
        document.add_paragraph().paragraph_format.space_after = Pt(1)

    document.add_paragraph("精准干预教学方案", style="Title")
    document.add_paragraph(value(basic["plan_title"]))
    document.add_paragraph(f'{basic["class_name"]} · {basic["subject"]} · {basic["total_duration_minutes"]} 分钟')
    field("审核结论", audit["overall_conclusion"])

    document.add_heading("1 精准教学基本信息", 1)
    table(["项目", "内容", "项目", "内容"], [
        ["学科", basic["subject"], "年级", basic["grade"]],
        ["教材版本与章节", basic["textbook_version_and_chapter"], "班级", basic["class_name"]],
        ["精准教学主题", basic["precision_teaching_topic"], "教师", basic["teacher_name"]],
    ], [2.4, 6.1, 2.4, 6.1])
    field("精准教学目标", basic["precision_teaching_goals"])
    field("精准教学内容", basic["precision_teaching_content"])

    document.add_heading("2 学生诊断结果", 1)
    field("诊断任务", diagnosis["diagnostic_task"]["task_text"])
    table(["SOLO 层级", "人数及比例", "典型表现", "主要进阶障碍"], [
        [f'{row["level"]} {row["level_name"]}', f'{row["count"]} 人 · {round(row["proportion"] * 100)}%',
         row.get("typical_performance"), row.get("main_obstacles")]
        for row in diagnosis["distribution_rows"]
    ], [2.5, 2.7, 5.9, 5.9])
    field("班级诊断总述", diagnosis["overall_summary"]["teacher_facing_paragraph"])

    document.add_heading("3 教学情境与条件", 1)
    for label, content in [
        ("教学内容与课标分析", context["teaching_content_and_curriculum_analysis"]),
        ("教学重点与难点分析", context["teaching_focus_and_difficulty_analysis"]),
        ("教学时长", f'{context["planned_duration_minutes"]} 分钟'),
        ("教学环境与 AI 支持条件", context["teaching_environment_and_ai_support_conditions"]),
    ]:
        field(label, content)

    document.add_heading("4 精准干预目标", 1)
    document.add_heading("共同核心目标", 2)
    document.add_paragraph(value(goals["common_core_goal"]["goal_statement"]))
    field("目标认知结构与理由", goals["common_core_goal"]["structure_and_rationale"])
    field("可观察的达成表现", goals["common_core_goal"]["observable_achievement"])
    document.add_heading("分层进阶目标", 2)
    table(["目标层级", "面向学生", "具体目标内容", "可观察的达成表现"], [
        [goal["target_level"], goal["students_display"], goal["goal_statement"], goal["observable_achievement"]]
        for goal in goals["progression_goals"]
    ], [2.6, 3.1, 5.65, 5.65])
    field("说明", goals["progression_goal_note"])

    document.add_heading("5 学习活动设计", 1)
    for index, activity in enumerate(plan["activities"], 1):
        document.add_heading(f'活动 {index} {activity["activity_name"]} {activity["duration_minutes"]} 分钟', 2)
        for label, item in [
            ("组织形式", activity["organization_forms"]),
            ("活动目标", activity["activity_objective"]),
            ("教师活动", activity["teacher_activities"]),
            ("学生活动", activity["student_activities"]),
            ("AI 辅助", activity.get("ai_support")),
            ("学习材料与资源", activity["learning_materials_and_resources"]),
            ("学习产出", activity["learning_product"]),
        ]:
            if item:
                field(label, item)

    document.add_heading("6 形成性评价", 1)
    for index, item in enumerate(plan["formative_evaluations"], 1):
        document.add_heading(f'评价 {index} {item["timing"]}', 2)
        for label, key in [
            ("评价目标", "evaluation_goal"),
            ("学习证据与判断标准", "evidence_and_criteria"),
            ("后续调整", "adjustment"),
            ("判断主体", "decision_by"),
        ]:
            field(label, item[key])

    document.add_heading("7 方案审核结果", 1)
    table(["检查项目", "状态", "发现的问题", "修改建议"], [
        [item["check_item"], item["status"], item["issue"], item["suggestion"]]
        for item in audit["review_items"]
    ], [3.1, 2.2, 5.85, 5.85])
    field("审核结论", audit["overall_conclusion"])

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
