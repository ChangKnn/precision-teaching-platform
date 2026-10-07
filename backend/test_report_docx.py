"""Editable Word export checks for the integrated teaching plan."""

import json
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException

from backend.app.intervention_integration_api import download_integration_report_docx
from backend.app.services.report_docx import render_report_docx


SAMPLE = Path(__file__).parent / "skills/precision-intervention-plan-integration-review/tests/sample-output.json"


class ReportDocxTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.result = json.loads(SAMPLE.read_text(encoding="utf-8"))

    def test_generates_editable_word_document(self):
        content = render_report_docx(self.result)
        self.assertTrue(content.startswith(b"PK"))
        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("精准干预教学方案", text)
        self.assertIn("形成性评价", text)
        self.assertGreaterEqual(len(document.tables), 3)
        self.assertIn("SOLO 层级", document.tables[1].cell(0, 0).text)
        document.paragraphs[0].text = "教师修改后的教学方案"
        edited = BytesIO()
        document.save(edited)
        self.assertEqual(Document(BytesIO(edited.getvalue())).paragraphs[0].text, "教师修改后的教学方案")

    @patch("backend.app.intervention_integration_api._source", return_value=({}, "new-hash", "1.0"))
    @patch("backend.app.intervention_integration_api._saved")
    def test_rejects_stale_report(self, saved, _source):
        saved.return_value = {"status": "stale", "result": self.result}
        with self.assertRaises(HTTPException) as caught:
            download_integration_report_docx("teaching", "classroom")
        self.assertEqual(caught.exception.status_code, 409)

    @patch("backend.app.intervention_integration_api._source", return_value=({}, "hash", "1.0"))
    @patch("backend.app.intervention_integration_api._saved")
    def test_download_is_word_attachment(self, saved, _source):
        saved.return_value = {"status": "current", "result": self.result}
        response = download_integration_report_docx("teaching", "classroom")
        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertTrue(response.body.startswith(b"PK"))
        self.assertIn("attachment", response.headers["content-disposition"])
        self.assertEqual(response.headers["cache-control"], "no-store")


if __name__ == "__main__":
    unittest.main()
